"""Golden-file regression test for the workbook writer.

Pins the *rendered output* of the deterministic pipeline — matcher -> report —
against a recorded baseline, so that a change to matching or to wording is
visible cell by cell instead of being discovered in a spreadsheet weeks later.

Why a golden file rather than assertions per behaviour: the thing at risk is the
workbook as a whole (values, colours, merges, block geometry), and enumerating
that by hand would be both enormous and less complete than a dump. A diff of the
JSON names the exact cell that moved.

WHAT THIS PROTECTS — see tests/README.md for the full list. In short: every
matcher verdict path, the per-paper block geometry, the status colours, the
merge ranges, and every user-visible string on all five sheets.

WHAT IT DOES NOT COVER: anything network-bound. ``resolve.py`` (Crossref DOI
resolution) and the Crossref branch of ``data_quality.py`` are excluded by
construction — this test calls :func:`match_paper` directly and supplies the
``data_quality`` field itself, so a run is fully offline and deterministic.

Run:
    python -m unittest discover -s tests -t . -v

Re-record the baseline after an intentional change (and explain the diff in the
commit message):
    python tests/test_report_golden.py --update
"""

import json
import re
import sys
import tempfile
import unittest
from pathlib import Path

_REPO = Path(__file__).resolve().parent.parent
_PKG = _REPO / "src" / "bealls_list_check"
# Same path shim the package's own entry scripts use (see bealls_list_check.py):
# the package root plus every foldered module (foo/foo.py), so the flat
# 'from x import y' imports resolve.
for _p in [_PKG, *(d for d in _PKG.iterdir() if d.is_dir() and (d / f"{d.name}.py").exists())]:
    if str(_p) not in sys.path:
        sys.path.insert(0, str(_p))
sys.path.insert(0, str(_PKG.parent))

import openpyxl                                          # noqa: E402
import report                                            # noqa: E402
from match import BeallIndex, match_paper                # noqa: E402

GOLDEN_PATH = Path(__file__).resolve().parent / "golden" / "report_baseline.json"

# The one volatile cell the writer emits: report.py's "Check run (local)" uses
# datetime.now().  Normalised so the dump is reproducible.
_TIMESTAMP_RE = re.compile(r"^\d{4}-\d{2}-\d{2} \d{2}:\d{2}:\d{2}$")


# ---------------------------------------------------------------------------
# Fixture: a synthetic snapshot + synthetic papers.
#
# Synthetic on purpose.  Real Beall entries would tie the baseline to the
# vendored snapshot, so refreshing that snapshot would break this test for a
# reason that has nothing to do with the code under test.  Every host here ends
# in .test, a reserved TLD (RFC 2606) that can never resolve.
# ---------------------------------------------------------------------------

SNAPSHOT_META = {
    "source": "https://example.test/fixture",
    "scraped_utc": "2026-01-01T00:00:00+00:00",
    "total_entries": 8,
    "note": "Synthetic fixture snapshot — not real data.",
}

ENTRIES = [
    {"name": "Dodgy Publishing Group", "url": "http://dodgypub.test/",
     "domain": "dodgypub.test", "normalized_name": "dodgy publishing group",
     "issns": [], "list_source": "publishers"},
    {"name": "Subhost Press", "url": "http://subhost.test/",
     "domain": "subhost.test", "normalized_name": "subhost press",
     "issns": [], "list_source": "publishers"},
    {"name": "Journal of Dubious Results", "url": "http://dubious.test/",
     "domain": "dubious.test", "normalized_name": "journal of dubious results",
     "issns": ["1111-2222"], "list_source": "standalone_journal"},
    {"name": "Journal of Shared Titles", "url": "http://sharedtitle.test/",
     "domain": "sharedtitle.test", "normalized_name": "journal of shared titles",
     "issns": [], "list_source": "standalone_journal"},
    # No website at all -> an exact name match is asserted (match.py:611).
    {"name": "Bulletin of Unlisted Sites", "url": "",
     "domain": "", "normalized_name": "bulletin of unlisted sites",
     "issns": [], "list_source": "standalone_journal"},
    {"name": "Cloned Medical Review", "url": "http://clone.test/",
     "domain": "clone.test", "normalized_name": "cloned medical review",
     "issns": [], "list_source": "hijacked",
     "authentic_url": "http://authentic.test/", "authentic_domain": "authentic.test"},
    {"name": "Vanity Books Ltd", "url": "http://vanitybooks.test/",
     "domain": "vanitybooks.test", "normalized_name": "vanity books ltd",
     "issns": [], "list_source": "vanity_press"},
    {"name": "Global Impact Factor Co", "url": "http://fakemetric.test/",
     "domain": "fakemetric.test", "normalized_name": "global impact factor co",
     "issns": [], "list_source": "misleading_metric"},
]

# (case label, expected status, paper record).  The label is echoed into the
# paper title so a failing cell in the workbook names its own case.
PAPERS = [
    ("domain -> on_list (publishers)", "on_list", {
        "paperId": "p01", "year": 2020, "externalIds": {"DOI": "10.1000/a"},
        "venue": "Some Journal",
        "publicationVenue": {"name": "Some Journal", "type": "journal",
                             "url": "http://dodgypub.test/j1"}}),
    ("domain_subdomain -> on_list", "on_list", {
        "paperId": "p02", "year": 2021, "externalIds": {"DOI": "10.1000/b"},
        "venue": "Sub Journal",
        "publicationVenue": {"name": "Sub Journal", "type": "journal",
                             "url": "http://ojs.subhost.test/x"}}),
    ("issn -> on_list (standalone)", "on_list", {
        "paperId": "p03", "year": 2022, "externalIds": {"DOI": "10.1000/c"},
        "venue": "Journal of Dubious Results",
        "publicationVenue": {"name": "Journal of Dubious Results",
                             "type": "journal", "issn": "1111-2222"}}),
    ("name_exact, listed entry has no site -> on_list", "on_list", {
        "paperId": "p04", "year": 2023, "externalIds": {"DOI": "10.1000/d"},
        "venue": "Bulletin of Unlisted Sites",
        "publicationVenue": {"name": "Bulletin of Unlisted Sites", "type": "journal",
                             "url": "http://somewhere-else.test/x"}}),
    ("name_exact uncorroborated -> review", "review", {
        "paperId": "p05", "year": 2023, "externalIds": {"DOI": "10.1000/e"},
        "venue": "Journal of Shared Titles",
        "publicationVenue": {"name": "Journal of Shared Titles", "type": "journal",
                             "url": "http://elsewhere.test/x"}}),
    ("hijacked by name -> review", "review", {
        "paperId": "p06", "year": 2022, "externalIds": {"DOI": "10.1000/f"},
        "venue": "Cloned Medical Review",
        "publicationVenue": {"name": "Cloned Medical Review", "type": "journal",
                             "url": "http://authentic.test/j"}}),
    ("vanity_press (weak) -> review", "review", {
        "paperId": "p07", "year": 2019, "externalIds": {"DOI": "10.1000/g"},
        "venue": "Vanity Books Ltd",
        "publicationVenue": {"name": "Vanity Books Ltd", "type": "journal",
                             "url": "http://vanitybooks.test/b"}}),
    ("misleading_metric (weak) -> review", "review", {
        "paperId": "p08", "year": 2019, "externalIds": {"DOI": "10.1000/h"},
        "venue": "Global Impact Factor Co",
        "publicationVenue": {"name": "Global Impact Factor Co", "type": "journal",
                             "url": "http://fakemetric.test/m"}}),
    ("name_fuzzy -> review", "review", {
        "paperId": "p09", "year": 2021, "externalIds": {"DOI": "10.1000/i"},
        "venue": "Journal of Dubius Results",
        "publicationVenue": {"name": "Journal of Dubius Results", "type": "journal"}}),
    # The multi-source case: two alternate URLs hitting entries on two DIFFERENT
    # lists, so list_source is '; '-joined.  This is the row that exercises the
    # joined-value handling in report.py and match.py.
    ("alternate URLs, two lists -> review (joined)", "review", {
        "paperId": "p10", "year": 2024, "externalIds": {"DOI": "10.1000/j"},
        "venue": "Merged Venue",
        "publicationVenue": {"name": "Merged Venue", "type": "journal",
                             "url": "http://innocent.test/j",
                             "alternate_urls": ["http://clone.test/a",
                                                "http://dodgypub.test/b"]}}),
    ("clean", "clean", {
        "paperId": "p11", "year": 2023, "externalIds": {"DOI": "10.1000/k"},
        "venue": "Respectable Journal",
        "publicationVenue": {"name": "Respectable Journal", "type": "journal",
                             "url": "http://respectable.test/j", "issn": "9999-8888"}}),
    ("preprint -> no_venue", "no_venue", {
        "paperId": "p12", "year": 2024, "externalIds": {"DOI": "10.1000/l"},
        "venue": "arXiv",
        "publicationVenue": {"name": "arXiv", "type": "repository",
                             "url": "https://arxiv.org/abs/1"}}),
    ("no venue information -> no_venue", "no_venue", {
        "paperId": "p13", "year": 2020, "externalIds": {},
        "venue": "", "publicationVenue": {}}),
    # alternate_names as an int makes list() raise inside _classify, which
    # match_paper turns into an 'error' row rather than killing the run.
    ("bad record -> error", "error", {
        "paperId": "p14", "year": 2020, "externalIds": {"DOI": "10.1000/n"},
        "venue": "Broken",
        "publicationVenue": {"name": "Broken", "type": "journal",
                             "alternate_names": 5}}),
]

# Fixed data-quality warnings on two papers, so the 'Data quality' sheet is
# rendered (report.py omits it entirely when nothing is flagged) without calling
# the network-dependent validator.
DATA_QUALITY = {
    "p10": ["Merged venue: this record points at 2 different publishers' sites.",
            "Fixture warning, second line."],
    "p13": ["No venue information in this record."],
}


def build_results():
    """Match every fixture paper and return the result dicts, in fixture order."""
    index = BeallIndex([dict(e) for e in ENTRIES])
    results = []
    for label, _expected, paper in PAPERS:
        paper = dict(paper, title=f"[{label}]")
        r = match_paper(index, paper, "fixture.json")
        r["data_quality"] = DATA_QUALITY.get(r["paperId"], [])
        results.append(r)
    return results


# ---------------------------------------------------------------------------
# Workbook -> canonical structure
# ---------------------------------------------------------------------------

def _cell_value(value):
    """Normalise a cell value: volatile timestamps become a fixed token."""
    if isinstance(value, str) and _TIMESTAMP_RE.match(value):
        return "<TIMESTAMP>"
    return value


def _fill_of(cell):
    """The cell's solid fill colour, or None. Kept because status colouring is
    a documented output (config.STATUS_FILL_COLORS) that a refactor could drop."""
    fill = cell.fill
    if not fill or fill.patternType != "solid":
        return None
    rgb = getattr(fill.fgColor, "rgb", None)
    return rgb if isinstance(rgb, str) else None


def dump_workbook(path):
    """Canonical, comparable structure for a workbook.

    Captures what a reader would notice changing: which sheets exist, their
    dimensions, every non-empty cell's value / fill / bold / hyperlink, the
    merge ranges (the per-paper block geometry), and the row heights and column
    widths (the layout).  Deliberately excludes anything openpyxl derives from
    wall-clock time.
    """
    wb = openpyxl.load_workbook(path)
    sheets = []
    for ws in wb.worksheets:
        cells = []
        for row in ws.iter_rows():
            for c in row:
                if c.value is None or c.value == "":
                    continue
                entry = {"ref": c.coordinate, "value": _cell_value(c.value)}
                fill = _fill_of(c)
                if fill:
                    entry["fill"] = fill
                if c.font is not None and c.font.bold:
                    entry["bold"] = True
                if c.hyperlink is not None:
                    entry["link"] = c.hyperlink.target or c.hyperlink.location
                cells.append(entry)
        sheets.append({
            "title": ws.title,
            "max_row": ws.max_row,
            "max_column": ws.max_column,
            "freeze_panes": ws.freeze_panes,
            "merged_ranges": sorted(str(r) for r in ws.merged_cells.ranges),
            "column_widths": {k: round(float(d.width), 2)
                              for k, d in sorted(ws.column_dimensions.items())
                              if d.width},
            "row_heights": {str(k): round(float(d.height), 2)
                            for k, d in sorted(ws.row_dimensions.items())
                            if d.height},
            "cells": cells,
        })
    return {"sheets": sheets}


def build_dump():
    """Run the fixture through the real writer and return the canonical dump."""
    results = build_results()
    with tempfile.TemporaryDirectory(prefix="bealls_golden_") as tmp:
        # report.py resolves its output folder from the module-level RESULTS_DIR
        # it imported from config; point that at a temp dir so the test never
        # writes into the package's results/ folder.
        original = report.RESULTS_DIR
        report.RESULTS_DIR = Path(tmp)
        try:
            out_path, n_flagged = report.save_workbook(
                results, dict(SNAPSHOT_META), ["fixture.json"],
                results_filename="golden_fixture.xlsx", n_excluded=3)
            dump = dump_workbook(out_path)
        finally:
            report.RESULTS_DIR = original
    dump["n_flagged"] = n_flagged
    dump["statuses"] = [r["status"] for r in results]
    dump["list_sources"] = [r["list_source"] for r in results]
    dump["matched_on"] = [r["matched_on"] for r in results]
    return dump


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------

class TestMatcherVerdicts(unittest.TestCase):
    """The verdict each fixture paper must reach.

    Stated separately from the golden file so that a matcher regression names
    the rule it broke, rather than showing up as an unexplained cell diff.
    """

    def test_each_case_reaches_its_documented_status(self):
        results = build_results()
        self.assertEqual(len(results), len(PAPERS))
        for (label, expected, _paper), r in zip(PAPERS, results):
            with self.subTest(case=label):
                self.assertEqual(r["status"], expected,
                                 f"{label}: expected {expected}, got {r['status']}")

    def test_all_five_list_categories_are_exercised(self):
        seen = set()
        for r in build_results():
            for part in (r["list_source"] or "").split(";"):
                if part.strip():
                    seen.add(part.strip())
        self.assertEqual(seen, {"publishers", "standalone_journal", "hijacked",
                                "vanity_press", "misleading_metric"})

    def test_joined_list_source_case_is_present(self):
        """The multi-source row must stay in the fixture — several bugs this
        suite guards against are only reachable through a joined value."""
        joined = [r for r in build_results() if ";" in (r["list_source"] or "")]
        self.assertEqual(len(joined), 1)
        self.assertEqual(joined[0]["list_source"], "hijacked; publishers")
        self.assertEqual(joined[0]["matched_on"], "domain_alternate_url")


class TestWorkbookGolden(unittest.TestCase):
    """Cell-level equivalence against the recorded baseline."""

    maxDiff = None

    def test_workbook_matches_golden(self):
        if not GOLDEN_PATH.exists():
            self.fail(f"No baseline at {GOLDEN_PATH}. Record one with:\n"
                      f"    python tests/test_report_golden.py --update")
        expected = json.loads(GOLDEN_PATH.read_text(encoding="utf-8"))
        actual = build_dump()

        # Compare sheet by sheet, then cell by cell, so a failure names the
        # sheet and the cell rather than printing two whole workbooks.
        self.assertEqual([s["title"] for s in actual["sheets"]],
                         [s["title"] for s in expected["sheets"]],
                         "the set or order of sheets changed")
        for got, want in zip(actual["sheets"], expected["sheets"]):
            with self.subTest(sheet=want["title"]):
                for key in ("max_row", "max_column", "freeze_panes",
                            "merged_ranges", "column_widths", "row_heights"):
                    self.assertEqual(got[key], want[key],
                                     f"{want['title']}: {key} changed")
                want_cells = {c["ref"]: c for c in want["cells"]}
                got_cells = {c["ref"]: c for c in got["cells"]}
                self.assertEqual(sorted(got_cells), sorted(want_cells),
                                 f"{want['title']}: the set of non-empty cells changed")
                for ref in sorted(want_cells):
                    self.assertEqual(got_cells[ref], want_cells[ref],
                                     f"{want['title']}!{ref} changed")

        for key in ("n_flagged", "statuses", "list_sources", "matched_on"):
            self.assertEqual(actual[key], expected[key], f"{key} changed")


def _update():
    GOLDEN_PATH.parent.mkdir(parents=True, exist_ok=True)
    GOLDEN_PATH.write_text(
        json.dumps(build_dump(), indent=1, ensure_ascii=False) + "\n",
        encoding="utf-8")
    print(f"Recorded baseline -> {GOLDEN_PATH}")


if __name__ == "__main__":
    if "--update" in sys.argv:
        _update()
    else:
        unittest.main()
